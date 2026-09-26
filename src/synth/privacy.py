"""Локальная проверка общих цепочек из 8 слов с реальным свободным текстом.

На экран выходят только числа и ID синтетических строк. Реальные значения
никогда не сохраняются и не печатаются. Запускать перед приёмкой корпуса.
"""
from __future__ import annotations

import hashlib
import json
import re
from collections import defaultdict
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from src.adapters.adapters import BASE_DIR
from src.synth.checks import load_raw

ROOT = Path(BASE_DIR)
WORD = re.compile(r"[^\W_]+", re.UNICODE)
WINDOW = 8
CSV_SOURCES = {
    "Обращения граждан 109 - Карагандинская область.csv": ("appeal_address",),
    "Обращения жителей 109 - Костанайская область.csv": ("result",),
    "Обращения граждан 109 - Восточно-Казахстанская область.csv": ("com_exp",),
    "Обращения граждан 109 - Алматинская область.csv": ("com_exp", "contractor"),
    "Обращения граждан 109 - Акмолинская область.csv": ("request_subject",),
}
XLSX_SOURCES = {
    "kostanay_109_incidents.xlsx": ("Обращения жителей 109 - Костанайская область.csv", ("result",)),
    "east_kazakhstan_109_appeals.xlsx": (
        "Обращения граждан 109 - Восточно-Казахстанская область.csv", ("com_exp",)),
    # Акмолинский xlsx — точный дубль csv, раздел 5; повторная проверка не нужна.
}


def signatures(text: str):
    words = WORD.findall(text.casefold())
    for i in range(len(words) - WINDOW + 1):
        payload = " ".join(words[i:i + WINDOW]).encode("utf-8")
        yield hashlib.blake2b(payload, digest_size=16).digest()


def _synthetic_rows() -> list[dict]:
    raw, duplicates = load_raw()
    if duplicates:
        raise ValueError(f"повторных id в A: {len(duplicates)}")
    rows = [{"id": i, "text": t} for i, t in raw.items()]
    c = Path("tests/fixtures/synth/templates_c.jsonl")
    if c.exists():
        rows += [json.loads(line) for line in c.read_text(encoding="utf-8").splitlines()
                 if line.strip()]
    return rows


def check(rows: list[dict] | None = None) -> dict:
    rows = _synthetic_rows() if rows is None else rows
    signatures_to_ids: dict[bytes, set[str]] = defaultdict(set)
    for row in rows:
        for sig in signatures(row["text"]):
            signatures_to_ids[sig].add(row["id"])
    hits: set[str] = set()
    scanned = 0

    def scan(value):
        nonlocal scanned
        if not isinstance(value, str) or not value.strip():
            return
        scanned += 1
        for sig in signatures(value):
            hits.update(signatures_to_ids.get(sig, ()))

    for name, fields in CSV_SOURCES.items():
        path = ROOT / name
        if not path.is_file():
            raise FileNotFoundError(path)
        headers = set(pd.read_csv(path, nrows=0).columns)
        absent = set(fields) - headers
        if absent:
            raise ValueError(f"{name}: нет колонок {sorted(absent)}")
        for chunk in pd.read_csv(path, dtype=str, usecols=list(fields), chunksize=20_000):
            for field in fields:
                for value in chunk[field].dropna():
                    scan(value)

    for name, (csv_name, fields) in XLSX_SOURCES.items():
        path = ROOT / name
        if not path.is_file():
            raise FileNotFoundError(path)
        headers = list(pd.read_csv(ROOT / csv_name, nrows=0).columns)
        indices = [headers.index(field) for field in fields]
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            for row in book.active.iter_rows(values_only=True):
                for idx in indices:
                    if idx < len(row):
                        scan(row[idx])
        finally:
            book.close()

    return {"synthetic_rows": len(rows), "real_values_scanned": scanned,
            "synthetic_ids_with_overlap": sorted(hits),
            "overlap_count": len(hits)}


def main() -> None:
    result = check()
    print(f"синтетических строк: {result['synthetic_rows']}; "
          f"реальных непустых значений проверено: {result['real_values_scanned']}; "
          f"совпадений из {WINDOW}+ слов: {result['overlap_count']}")
    if result["overlap_count"]:
        print("ID для переписывания:", ", ".join(result["synthetic_ids_with_overlap"][:50]))
        raise SystemExit(1)


if __name__ == "__main__":
    main()
