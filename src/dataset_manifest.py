"""Версия raw: безопасные агрегаты, явное принятие, проверка до сборки (0e)."""
from __future__ import annotations

import argparse
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import sys

import pandas as pd
from openpyxl import load_workbook

from src import paths
from src.adapters.schemas import FILES, HEAD

DATES = {"kar": ("created_date", "%m/%d/%y %H:%M"),
         "kos": ("createddate", "ISO8601"), "tur": ("createddate", "ISO8601"),
         "vko": ("creation_date", "%d.%m.%Y %H:%M:%S"),
         "alm": ("creation_date", "%d.%m.%Y %H:%M:%S"),
         "akm": ("creation_date", "ISO8601"), "pav": ("create_date", "ISO8601")}
XLSX = {"akmola_109_appeals.xlsx": "akm", "east_kazakhstan_109_appeals.xlsx": "vko",
        "kostanay_109_incidents.xlsx": "kos", "turkestan_109_incidents (1).xlsx": "tur"}
ACCEPTED = paths.DATA_DIR / "raw_manifest.json"
CANDIDATE = paths.DATA_DIR / "raw_candidate.json"


class ManifestError(ValueError):
    """Сообщение состоит только из контролируемых строк, без исходных значений."""


def digest(path):
    with Path(path).open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def version(files):
    payload = json.dumps(files, sort_keys=True, ensure_ascii=False, separators=(",", ":"))
    return hashlib.sha256(payload.encode()).hexdigest()


def expected(include_xlsx=True):
    entries = [(name, key[:3], False) for key, name in FILES.items()]
    if include_xlsx:
        entries += [(name, family, True) for name, family in XLSX.items()]
    return sorted(entries)


def date_summary(chunks, col, fmt):
    count = failures = 0
    lower = upper = None
    for values in chunks:
        series = pd.Series(values)
        parsed = pd.to_datetime(series, format=fmt, errors="coerce")
        count += len(parsed)
        failures += int(parsed.isna().sum())
        valid = parsed.dropna()
        if len(valid):
            a, b = valid.min(), valid.max()
            lower = a if lower is None else min(lower, a)
            upper = b if upper is None else max(upper, b)
    if not count or failures / count > 0.01:
        raise ManifestError("Пустой файл или непарсящихся дат больше 1%; нужна проверка схемы.")
    return {"rows": count, "date_failures": failures,
            "first_date": lower.date().isoformat(), "last_date": upper.date().isoformat()}


def xlsx_dates(path, family):
    columns = HEAD[family].split(",")
    index = columns.index(DATES[family][0])
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        if len(book.worksheets) != 1:
            raise ManifestError("XLSX: требуется один лист.")
        sheet = book.active
        if sheet.max_column != len(columns):
            raise ManifestError("XLSX: несовместимая ширина схемы.")
        rows = sheet.iter_rows(values_only=True)
        if family == "akm" and list(next(rows, ())) != columns:
            raise ManifestError("XLSX: несовместимый заголовок.")
        batch = []
        for row in rows:
            batch.append(row[index])
            if len(batch) == 20000:
                yield batch
                batch = []
        if batch:
            yield batch
    finally:
        book.close()


def generate(raw, *, include_xlsx=True, source="real"):
    files = []
    for name, family, excel in expected(include_xlsx):
        path = Path(raw) / name
        if not path.is_file():
            raise ManifestError(f"НЕ ХВАТАЕТ ожидаемого файла: {name}")
        try:
            before = digest(path)
            col, fmt = DATES[family]
            if excel:
                chunks = xlsx_dates(path, family)
            else:
                if list(pd.read_csv(path, nrows=0).columns) != HEAD[family].split(","):
                    raise ManifestError("CSV: несовместимый заголовок; обновите контракт адаптера.")
                chunks = (frame[col] for frame in pd.read_csv(
                    path, dtype=str, usecols=[col], chunksize=20000))
            entry = {"relative_path": name, "size_bytes": path.stat().st_size,
                     "sha256": before, "schema_family": family,
                     "schema_status": "structurally-compatible" if excel else "compatible"}
            entry.update(date_summary(chunks, col, fmt))
            if digest(path) != before:
                raise ManifestError("Файл менялся во время проверки; повторите на снимке сырья.")
            files.append(entry)
        except ManifestError:
            raise
        except Exception:
            # Текст parser exception может включать сырую строку и ПДн.
            raise ManifestError(f"Не удалось безопасно прочитать ожидаемый файл: {name}") from None
    return {"format_version": 1, "source": source, "dataset_id": version(files),
            "created_at": datetime.now(timezone.utc).isoformat(), "files": files}


def ensure(condition):
    if not condition:
        raise ValueError("invalid manifest")


def read_manifest(path, *, source="real", include_xlsx=True):
    try:
        obj = json.loads(Path(path).read_text(encoding="utf-8"))
        ensure(set(obj) == {"format_version", "source", "dataset_id", "created_at", "files"})
        datetime.fromisoformat(obj["created_at"])
        ensure(obj["format_version"] == 1 and obj["source"] == source)
        files = obj["files"]
        ensure(obj["dataset_id"] == version(files))
        ensure([f["relative_path"] for f in files] == [x[0] for x in expected(include_xlsx)])
        for entry, (_, family, excel) in zip(files, expected(include_xlsx)):
            ensure(set(entry) == {"relative_path", "size_bytes", "sha256", "schema_family",
                                  "schema_status", "rows", "date_failures", "first_date", "last_date"})
            ensure(entry["schema_family"] == family)
            ensure(entry["schema_status"] == ("structurally-compatible" if excel else "compatible"))
            ensure(type(entry["rows"]) is int and entry["rows"] > 0)
            ensure(type(entry["size_bytes"]) is int and entry["size_bytes"] > 0)
            ensure(len(entry["sha256"]) == 64 and set(entry["sha256"]) <= set("0123456789abcdef"))
            ensure(0 <= entry["date_failures"] / entry["rows"] <= .01)
            ensure(datetime.fromisoformat(entry["first_date"]) <= datetime.fromisoformat(entry["last_date"]))
        return obj
    except Exception:
        raise ManifestError("Принятый manifest отсутствует или повреждён. "
                            "Создайте кандидат: python -m src.dataset_manifest; "
                            "после просмотра примите: python -m src.dataset_manifest --accept ID.") from None


def verify(raw=None, manifest=None, *, source="real", include_xlsx=True):
    raw = paths.RAW_DIR if raw is None else Path(raw)
    obj = read_manifest(ACCEPTED if manifest is None else manifest,
                        source=source, include_xlsx=include_xlsx)
    changed = []
    for entry in obj["files"]:
        path = raw / entry["relative_path"]
        try:
            same = path.stat().st_size == entry["size_bytes"] and digest(path) == entry["sha256"]
        except OSError:
            same = False
        if not same:
            changed.append(entry["relative_path"])
    if changed:
        raise ManifestError("Выгрузка отличается от ранее зафиксированной: " + "; ".join(changed)
                            + ". Сборка не начата. Создайте и просмотрите кандидат "
                            "python -m src.dataset_manifest, затем --accept ID.")
    return obj


def write_json(path, obj):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(obj, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    temporary.replace(path)


def accept(raw, candidate, identity, target=ACCEPTED, *, source="real", include_xlsx=True):
    obj = read_manifest(candidate, source=source, include_xlsx=include_xlsx)
    if identity != obj["dataset_id"]:
        raise ManifestError("ID не совпадает с просмотренным кандидатом; принятие отменено.")
    verify(raw, candidate, source=source, include_xlsx=include_xlsx)
    target = Path(target)
    archive = target.parent / "manifests" / (identity + ".json")
    if not archive.exists():
        write_json(archive, obj)
    write_json(target, obj)
    return obj


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    group = parser.add_mutually_exclusive_group()
    group.add_argument("--accept", metavar="ID")
    group.add_argument("--check", action="store_true")
    args = parser.parse_args()
    try:
        if paths.FAKE:
            raise ManifestError("Команда принятия предназначена для REAL; fixture FAKE имеет свой manifest.")
        paths.validate_raw_source(paths.RAW_DIR, recursive=True)
        if args.check:
            obj = verify()
            print("Manifest совпадает с принятой версией:", obj["dataset_id"])
        elif args.accept:
            obj = accept(paths.RAW_DIR, CANDIDATE, args.accept)
            print("Принята версия:", obj["dataset_id"])
        else:
            obj = generate(paths.RAW_DIR)
            write_json(CANDIDATE, obj)
            print("Кандидат data/raw_candidate.json; файлов:", len(obj["files"]))
            print("dataset_id:", obj["dataset_id"])
            print("Просмотрите агрегаты, затем: python -m src.dataset_manifest --accept ID")
        return 0
    except ManifestError as exc:
        print(str(exc), file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
