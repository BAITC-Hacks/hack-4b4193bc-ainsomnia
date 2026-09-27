#!/usr/bin/env python3
"""Выгрузка витрины в Excel и PDF.

Собирает ровно то, что видно на экране при текущих фильтрах: ничего не
досчитывает и не пересчитывает. Таблицы приходят готовыми из `src.dashboard`.

Строка ограничений НЕ сочиняется здесь: она вытаскивается из CLAUDE.md по
якорным фразам. Если формулировку в CLAUDE.md переименуют, сборка упадёт с
явной ошибкой — это лучше, чем молча выгрузить отчёт без оговорки.

Персональных данных в выгрузке нет по построению: канон их не содержит
(раздел 1 CLAUDE.md). Проверка готового файла — `check_export_pii`.
"""
from __future__ import annotations

import io
import os
import re
from datetime import date
from pathlib import Path

import pandas as pd

CLAUDE_MD = Path("CLAUDE.md")
# Якоря в CLAUDE.md: начало строки -> сама строка целиком идёт в отчёт.
LIMIT_ANCHORS = (
    "**Регионы применяют разные определения обращения.",
    "2. Любая сводка или график, где регионы сравниваются между собой",
)
PII_COLUMNS = ("full_name", "applicant_number", "operator", "xcoordinate",
               "ycoordinate", "appeal_address", "street", "com_exp", "result",
               "request_subject")
# Подписи классов в выгрузке — те же, что на экране витрины.
CLASS_COLS = {"problem": "жалобы на городские проблемы", "info": "справочные звонки",
              "system": "служебные записи"}
# Префикс +7, 7 или 8; код оператора в скобках или без; группы через пробел, дефис или слитно.
# До 2026-09-27 префикс 8 не ловился вовсе: «8 700 000 00 00», «8-700-000-00-00» проходили
# мимо всех проверок ПДн (CLAUDE.md, раздел 1). Границы по цифрам — чтобы 12 цифр подряд
# (ИИН) ловил LONG_DIGITS, а не телефон внутри них.
PHONE = re.compile(r"(?<!\d)(?:\+?7|8)[\s\-]*\(?\s*\d{3}\s*\)?[\s\-]*\d{3}[\s\-]*\d{2}[\s\-]*\d{2}(?!\d)")
LONG_DIGITS = re.compile(r"\d{12}")
ADDR = re.compile(r"(?:\bдом\b.*\bкв\b)|(?:^|\s)ул\.|(?:^|\s)мкр\b", re.I)


# ---------------------------------------------------------------- ограничения
def read_limits(path=CLAUDE_MD):
    """Достать формулировки об ограничениях из CLAUDE.md по якорям."""
    text = path.read_text(encoding="utf-8")
    out = []
    for anchor in LIMIT_ANCHORS:
        found = [ln.strip() for ln in text.splitlines() if ln.strip().startswith(anchor)]
        if not found:
            raise ValueError(
                f"В {path} не найдена строка, начинающаяся с «{anchor[:50]}…». "
                "Текст ограничений берётся из CLAUDE.md и не сочиняется в коде: "
                "поправьте якорь в LIMIT_ANCHORS или верните формулировку.")
        # снять markdown-разметку и номер пункта списка из CLAUDE.md
        out.append(re.sub(r"^\d+\.\s*", "", re.sub(r"[*`]", "", found[0])))
    return out


def limits_block(regions, period, filters=(), path=CLAUDE_MD):
    """Строки ограничений для первой страницы PDF и первого листа Excel.

    `filters` — человекочитаемое описание того, что было выбрано на экране.
    Выгрузка повторяет экран, и по файлу должно быть видно, при каких фильтрах
    он снят, иначе числа нельзя сопоставить ни с чем."""
    lo, hi = period
    return [
        f"Регионы в выгрузке: {', '.join(regions)} — {len(regions)} шт.",
        f"Период: {lo:%d.%m.%Y} — {hi:%d.%m.%Y}.",
        *[f"Фильтр: {f}" for f in filters],
        *read_limits(path),
        f"Выгружено {date.today():%d.%m.%Y}. Источник — сводная таблица обращений "
        f"109 (файл data/unified.parquet). "
        f"Числа не пересчитывались: выгружено то же, что показано на экране.",
    ]


# ---------------------------------------------------------------- таблицы
def sheet_regions(df):
    """Сводка по регионам: всего и по классам обращения."""
    g = (df.pivot_table(index="region", columns="appeal_class", values="created_at",
                        aggfunc="size", fill_value=0)
           .rename(columns=CLASS_COLS))
    cols = list(CLASS_COLS.values())
    for c in cols:
        if c not in g.columns:
            g[c] = 0
    g = g[cols]
    g.insert(0, "всего", g.sum(axis=1))
    # доля info + system: «не жалобы» — подпись для руководителя, смысл прежний
    g["не жалобы, % от всех"] = ((g["всего"] - g[CLASS_COLS["problem"]])
                                 / g["всего"].where(g["всего"] > 0) * 100).round(1)
    return g.sort_values("всего", ascending=False).reset_index().rename(
        columns={"region": "регион"})


def sheet_topics(df):
    """Структура тем внутри класса problem."""
    p = df[df.appeal_class == "problem"]
    g = p.groupby("topic").size().rename("обращений").sort_values(ascending=False)
    out = g.reset_index().rename(columns={"topic": "тема"})
    total = int(g.sum())
    out["доля, %"] = (out["обращений"] / total * 100).round(2) if total else 0.0
    return out


def export_frames(df, events):
    """Единственный источник содержимого для обоих форматов.

    Excel и PDF обязаны строиться из одних и тех же таблиц, иначе проверка одного
    файла ничего не говорит про другой."""
    return {"Сводка по регионам": sheet_regions(df),
            "Темы жалоб": sheet_topics(df),
            "Всплески жалоб": events}


# ---------------------------------------------------------------- Excel
def build_excel(df, events, regions, period, filters=(), path=CLAUDE_MD):
    """Три листа плюс лист ограничений. Возвращает bytes."""
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        pd.DataFrame({"Ограничения выгрузки":
                      limits_block(regions, period, filters, path)}) \
            .to_excel(w, sheet_name="Ограничения", index=False)
        for name, frame in export_frames(df, events).items():
            frame.to_excel(w, sheet_name=name, index=False)
        for name, width in (("Ограничения", 120), ("Сводка по регионам", 22),
                            ("Темы жалоб", 34), ("Всплески жалоб", 30)):
            ws = w.book[name]
            ws.column_dimensions["A"].width = width
    return buf.getvalue()


# ---------------------------------------------------------------- PDF
def _pdf_table(data, styles, col_widths=None):
    from reportlab.lib import colors
    from reportlab.platypus import Table, TableStyle
    t = Table(data, colWidths=col_widths, repeatRows=1)
    t.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), styles["font"]),
        ("FONTNAME", (0, 0), (-1, 0), styles["bold"]),
        ("FONTSIZE", (0, 0), (-1, -1), 7.5),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#ECEFF1")),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#B0BEC5")),
        ("ALIGN", (1, 1), (-1, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 2),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
    ]))
    return t


class PdfUnavailable(RuntimeError):
    """PDF собрать нельзя по внешней причине — текст говорит, что сделать.
    Excel от этих причин не зависит."""


# Шрифты с кириллицей по порядку: свой через переменную, Linux (Debian/Ubuntu и
# образ контейнера, Fedora), macOS. Первым найденным и печатаем.
FONT_CANDIDATES = [
    ("Custom", os.environ.get("NAZAR_PDF_FONT", ""), os.environ.get("NAZAR_PDF_FONT_BOLD")),
    ("DejaVuSans", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
     "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
    ("DejaVuSans", "/usr/share/fonts/dejavu-sans-fonts/DejaVuSans.ttf",
     "/usr/share/fonts/dejavu-sans-fonts/DejaVuSans-Bold.ttf"),
    ("DejaVuSans", "/opt/homebrew/share/fonts/DejaVuSans.ttf", None),
    ("Arial", "/System/Library/Fonts/Supplemental/Arial.ttf",
     "/System/Library/Fonts/Supplemental/Arial Bold.ttf"),
    ("ArialUnicode", "/System/Library/Fonts/Supplemental/Arial Unicode.ttf", None),
]


def _register_font():
    """Шрифт с кириллицей. Без него reportlab печатает чёрные квадраты."""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    for name, regular, bold in FONT_CANDIDATES:
        if regular and Path(regular).exists():
            pdfmetrics.registerFont(TTFont(name, regular))
            if bold and Path(bold).exists():
                pdfmetrics.registerFont(TTFont(name + "-Bold", bold))
                return {"font": name, "bold": name + "-Bold"}
            return {"font": name, "bold": name}
    raise PdfUnavailable(
        "не найден шрифт с кириллицей — без него текст в PDF вышел бы квадратами. "
        "Что сделать: на Debian/Ubuntu — `apt install fonts-dejavu-core`; на macOS "
        "подходящий Arial есть в системе; либо укажите свой TTF в переменной "
        "NAZAR_PDF_FONT=/путь/к/шрифту.ttf. Искали: "
        + "; ".join(r for _, r, _ in FONT_CANDIDATES if r))


CHART_HINT = ("для картинок графиков нужен браузер Chrome или Chromium — его "
              "использует пакет kaleido. Установите Chrome или выполните "
              "`.venv/bin/plotly_get_chrome`")


def _fig_png(fig, width=980, height=420):
    """Картинка графика для PDF: (png, None) или (None, причина).

    Раньше любая ошибка глоталась, и PDF молча выходил без графиков."""
    try:
        return fig.to_image(format="png", width=width, height=height, scale=2), None
    except Exception as e:                       # noqa: BLE001 — причина уходит в отчёт
        return None, f"{CHART_HINT} (ошибка: {type(e).__name__})"


def build_pdf(df, events, regions, period, figures=(), filters=(), path=CLAUDE_MD,
              skipped=None):
    """Печатный отчёт: ограничения, две таблицы, лента, графики. Возвращает bytes.

    Не собирается — PdfUnavailable с объяснением (нет reportlab, нет шрифта).
    График не отрисовался — в PDF вместо него строка с причиной, а пара
    (название, причина) добавляется в `skipped`, чтобы вызывающий сказал об этом."""
    try:
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import (Image, PageBreak, Paragraph, SimpleDocTemplate,
                                        Spacer)
    except ImportError as e:
        raise PdfUnavailable(
            "не установлен пакет reportlab. Что сделать: переустановите зависимости "
            "проекта — `uv pip install --python .venv/bin/python -e .` из корня "
            "репозитория (reportlab указан в requirements.txt)") from e

    fonts = _register_font()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4),
                            leftMargin=14 * mm, rightMargin=14 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm,
                            title="Обращения в 109 — обзор для руководителя")
    base = getSampleStyleSheet()["Normal"]
    h1 = ParagraphStyle("h1", parent=base, fontName=fonts["bold"], fontSize=16,
                        leading=20, spaceAfter=8)
    h2 = ParagraphStyle("h2", parent=base, fontName=fonts["bold"], fontSize=11,
                        leading=14, spaceBefore=10, spaceAfter=5)
    txt = ParagraphStyle("txt", parent=base, fontName=fonts["font"], fontSize=8.5,
                         leading=12)

    story = [Paragraph("Обращения в 109 — обзор для руководителя", h1),
             Paragraph("Ограничения выгрузки", h2)]
    for line in limits_block(regions, period, filters, path):
        story.append(Paragraph("— " + line, txt))

    frames = export_frames(df, events)
    reg, top = frames["Сводка по регионам"], frames["Темы жалоб"]
    story += [Paragraph("Сводка по регионам", h2)]
    story.append(_pdf_table([list(reg.columns)] + reg.astype(str).values.tolist(), fonts))

    story += [Paragraph("Темы жалоб на городские проблемы", h2)]
    story.append(_pdf_table([list(top.columns)] + top.astype(str).values.tolist(), fonts))

    story += [PageBreak(), Paragraph(f"Всплески жалоб — {len(events)}", h2)]
    if len(events):
        ev = events.head(60)
        story.append(_pdf_table([list(ev.columns)] + ev.astype(str).values.tolist(), fonts))
        if len(events) > len(ev):
            story.append(Spacer(1, 4))
            story.append(Paragraph(
                f"Показаны первые {len(ev)} всплесков из {len(events)} в порядке "
                f"сортировки на экране. Полный список — на листе «Всплески жалоб» "
                f"в Excel-выгрузке.", txt))
    else:
        story.append(Paragraph("Под выбранные фильтры не попало ни одного всплеска.",
                               txt))

    for title, fig in figures:
        png, why = _fig_png(fig)
        story += [PageBreak(), Paragraph(title, h2)]
        if png is None:
            story.append(Paragraph(f"График не вставлен: {why}.", txt))
            if skipped is not None:
                skipped.append((title, why))
            continue
        story.append(Image(io.BytesIO(png), width=250 * mm, height=107 * mm))
    doc.build(story)
    return buf.getvalue()


# ---------------------------------------------------------------- проверка ПДн
def check_frames_pii(frames):
    """Проверка содержимого, общего для Excel и PDF: сами таблицы и их заголовки.

    PDF нельзя прочитать обратно так же просто, как xlsx (текст в нём закодирован
    подмножеством шрифта), поэтому источник проверяется до отрисовки. Оба формата
    строятся из `export_frames`, поэтому проверка покрывает оба."""
    hits = {"телефон": [], "12 цифр": [], "адрес": [], "колонка ПДн": []}
    cells = 0
    for name, frame in frames.items():
        for col in frame.columns:
            if any(c in str(col).lower() for c in PII_COLUMNS):
                hits["колонка ПДн"].append((name, str(col)))
        for col in frame.columns:
            ser = frame[col]
            if ser.dtype != object and not str(ser.dtype).startswith("str"):
                cells += len(ser)
                continue
            for v in ser.astype(str):
                cells += 1
                if any(c in v.lower() for c in PII_COLUMNS):
                    hits["колонка ПДн"].append((name, v[:2] + "*" * 8))
                if PHONE.search(v):
                    hits["телефон"].append((name, v[:2] + "*" * 8))
                if LONG_DIGITS.search(v):
                    hits["12 цифр"].append((name, v[:2] + "*" * 8))
                if ADDR.search(v):
                    hits["адрес"].append((name, v[:2] + "*" * 8))
    return {"таблиц": len(frames), "ячеек": cells,
            "нарушения": {k: v for k, v in hits.items() if v}}


def check_export_pii(xlsx_bytes):
    """Проверка ГОТОВОГО файла, а не намерений: читаем обратно и ищем ПДн.

    Возвращает словарь с числами по каждой проверке. Пустой список нарушений —
    единственный приемлемый результат.
    """
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(xlsx_bytes), read_only=True, data_only=True)
    cells = scanned = 0
    hits = {"телефон": [], "12 цифр": [], "адрес": [], "колонка ПДн": []}
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for v in row:
                cells += 1
                if not isinstance(v, str):
                    continue
                scanned += 1
                low = v.lower()
                for col in PII_COLUMNS:
                    if col in low:
                        hits["колонка ПДн"].append((ws.title, col))
                if PHONE.search(v):
                    hits["телефон"].append((ws.title, v[:2] + "*" * 8))
                if LONG_DIGITS.search(v):
                    hits["12 цифр"].append((ws.title, v[:2] + "*" * 8))
                if ADDR.search(v):
                    hits["адрес"].append((ws.title, v[:2] + "*" * 8))
    wb.close()
    return {"листов": len(wb.worksheets), "ячеек": cells, "текстовых": scanned,
            "нарушения": {k: v for k, v in hits.items() if v}}
